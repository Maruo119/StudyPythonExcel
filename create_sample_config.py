"""
サンプル設定Excelを生成するスクリプト
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

wb = Workbook()
wb.remove(wb.active)

# ============================================
# シート1: アプリマスタ
# ============================================
ws_apps = wb.create_sheet("アプリ")
ws_apps["A1"] = "アプリID"
ws_apps["B1"] = "アプリ名"

apps_data = [
    ("APP001", "ユーザー管理システム"),
    ("APP002", "決済システム"),
    ("APP003", "レポート生成システム"),
    ("APP004", "通知サービス"),
    ("APP005", "認証サービス"),
    ("APP006", "ログ管理システム"),
    ("APP007", "バックアップサービス"),
    ("APP008", "パフォーマンス監視"),
    ("APP009", "キャッシュシステム"),
    ("APP010", "データ同期エンジン"),
]

for idx, (app_id, app_name) in enumerate(apps_data, start=2):
    ws_apps[f"A{idx}"] = app_id
    ws_apps[f"B{idx}"] = app_name

# 列幅調整
ws_apps.column_dimensions["A"].width = 12
ws_apps.column_dimensions["B"].width = 20

# ============================================
# シート2: フェーズマスタ
# ============================================
ws_phases = wb.create_sheet("フェーズ")
ws_phases["A1"] = "フェーズID"
ws_phases["B1"] = "フェーズ名"
ws_phases["C1"] = "区分"

phases_data = [
    ("P001", "要件定義", "資産"),
    ("P002", "基本設計", "資産"),
    ("P003", "詳細設計", "資産"),
    ("P004", "実装", "費用"),
    ("P005", "単体テスト", "費用"),
    ("P006", "統合テスト", "費用"),
    ("P007", "本番環境構築", "費用"),
]

for idx, (phase_id, phase_name, division) in enumerate(phases_data, start=2):
    ws_phases[f"A{idx}"] = phase_id
    ws_phases[f"B{idx}"] = phase_name
    ws_phases[f"C{idx}"] = division

ws_phases.column_dimensions["A"].width = 12
ws_phases.column_dimensions["B"].width = 18
ws_phases.column_dimensions["C"].width = 10

# ============================================
# シート3: 工数データ
# ============================================
ws_work = wb.create_sheet("工数")
ws_work["A1"] = "アプリID"
ws_work["B1"] = "フェーズID"
ws_work["C1"] = "工数（人日）"

# サンプル工数データ（ランダムに配置）
work_data = [
    ("APP001", "P001", 5),
    ("APP001", "P002", 8),
    ("APP001", "P004", 15),
    ("APP001", "P005", 10),
    ("APP002", "P001", 10),
    ("APP002", "P002", 12),
    ("APP002", "P003", 20),
    ("APP002", "P004", 30),
    ("APP002", "P006", 18),
    ("APP003", "P002", 6),
    ("APP003", "P003", 10),
    ("APP003", "P004", 20),
    ("APP004", "P001", 3),
    ("APP004", "P002", 5),
    ("APP004", "P004", 12),
    ("APP004", "P005", 8),
    ("APP005", "P001", 4),
    ("APP005", "P003", 15),
    ("APP005", "P004", 25),
    ("APP006", "P002", 8),
    ("APP006", "P004", 18),
    ("APP006", "P005", 10),
    ("APP007", "P001", 3),
    ("APP007", "P003", 12),
    ("APP007", "P004", 16),
    ("APP008", "P002", 7),
    ("APP008", "P004", 14),
    ("APP009", "P003", 11),
    ("APP009", "P004", 22),
    ("APP010", "P001", 6),
    ("APP010", "P002", 9),
    ("APP010", "P004", 20),
    ("APP010", "P006", 12),
]

for idx, (app_id, phase_id, hours) in enumerate(work_data, start=2):
    ws_work[f"A{idx}"] = app_id
    ws_work[f"B{idx}"] = phase_id
    ws_work[f"C{idx}"] = hours

ws_work.column_dimensions["A"].width = 12
ws_work.column_dimensions["B"].width = 12
ws_work.column_dimensions["C"].width = 15

# ============================================
# ヘッダー行のスタイル統一
# ============================================
header_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
header_font = Font(bold=True, size=11)
header_align = Alignment(horizontal="center", vertical="center")

for ws in [ws_apps, ws_phases, ws_work]:
    for cell in ws[1]:
        if cell.value:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

# ファイルを保存
Path("input").mkdir(exist_ok=True)
wb.save("input\\config.xlsx")
print("[OK] サンプル設定ファイルを作成しました: input\\config.xlsx")
